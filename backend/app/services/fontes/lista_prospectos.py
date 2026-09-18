"""Leitura da lista de prospectos (extração da Receita por DDD, em .xlsx ou .csv).

Formato esperado (uma empresa por linha, cabeçalho na primeira linha):

    cnpj, razao_social, capital_social, porte, situacao_cadastral, cnae_fiscal,
    tipo_logradouro, logradouro, numero, bairro, cep, uf, ddd1,
    telefone_completo, correio_eletronico, nome_municipio

Só `cnpj` e `razao_social` são obrigatórios; o resto é aproveitado quando
existe. Os nomes das colunas são tolerantes (sem acento, maiúsculas ou
minúsculas, alguns apelidos), porque a lista muda de mãos e ninguém deveria
ter que reescrever o cabeçalho para importar.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass

from ..cnpj import normalizar_cnpj


class ListaInvalida(Exception):
    """Arquivo que não é uma lista de prospectos, ou está vazio."""


# nome canônico -> apelidos aceitos (já sem acento, minúsculo, com "_")
_APELIDOS = {
    "cnpj": {"cnpj", "cnpj_completo"},
    "razao_social": {"razao_social", "razao", "nome", "nome_empresarial", "empresa"},
    "capital_social": {"capital_social", "capital"},
    "porte": {"porte", "porte_empresa"},
    "situacao": {"situacao_cadastral", "situacao"},
    "cnae": {"cnae_fiscal", "cnae", "cnae_principal", "cnae_fiscal_principal"},
    "tipo_logradouro": {"tipo_logradouro"},
    "logradouro": {"logradouro", "endereco"},
    "numero": {"numero"},
    "bairro": {"bairro"},
    "cep": {"cep"},
    "uf": {"uf", "estado"},
    "ddd": {"ddd1", "ddd"},
    "telefone": {"telefone_completo", "telefone", "telefone1", "telefone_1", "fone"},
    "email": {"correio_eletronico", "email", "e_mail"},
    "cidade": {"nome_municipio", "municipio", "cidade"},
}

_COLUNAS_OBRIGATORIAS = ("cnpj", "razao_social")


@dataclass
class LinhaLista:
    cnpj: str
    razao_social: str
    capital_social: float | None
    porte: str
    situacao: str
    cnae: str
    logradouro: str
    numero: str
    bairro: str
    cep: str
    cidade: str
    uf: str
    telefone: str
    email: str


@dataclass
class ListaLida:
    linhas: list[LinhaLista]
    invalidas: int          # linhas sem CNPJ de 14 dígitos ou sem nome
    repetidas: int          # CNPJ que apareceu mais de uma vez no arquivo


def _chave(cabecalho) -> str:
    base = unicodedata.normalize("NFD", str(cabecalho or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", base.strip().lower()).strip("_")


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _numero(valor) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto:  # padrão brasileiro: 1.234,56
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _cnpj(valor) -> str:
    """Só dígitos, com 14 posições. O Excel guarda CNPJ como número e perde os
    zeros à esquerda ("00072266000140" vira 72266000140)."""
    if isinstance(valor, float):
        valor = int(valor)
    digitos = normalizar_cnpj(str(valor if valor is not None else ""))
    # o Excel só perde os zeros à esquerda (poucos); abaixo de 10 dígitos é lixo, não CNPJ
    if len(digitos) < 10 or len(digitos) > 14:
        return ""
    digitos = digitos.zfill(14)
    return "" if set(digitos) == {"0"} else digitos


def _linhas_brutas(conteudo: bytes) -> list[list]:
    """Devolve a tabela como lista de linhas (xlsx ou csv)."""
    if conteudo[:2] == b"PK":  # .xlsx é um zip
        from openpyxl import load_workbook
        try:
            planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as erro:
            raise ListaInvalida("Não consegui abrir a planilha. Ela está corrompida ou protegida por senha?") from erro
        aba = planilha.active
        return [list(linha) for linha in aba.iter_rows(values_only=True)]

    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1")
    amostra = texto[:4096]
    delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def ler(conteudo: bytes) -> ListaLida:
    """Converte a planilha em linhas normalizadas, sem repetir CNPJ."""
    tabela = _linhas_brutas(conteudo)
    if not tabela:
        raise ListaInvalida("O arquivo está vazio.")

    cabecalho = [_chave(c) for c in tabela[0]]
    posicao: dict[str, int] = {}
    for canonico, apelidos in _APELIDOS.items():
        for i, nome in enumerate(cabecalho):
            if nome in apelidos:
                posicao[canonico] = i
                break

    faltando = [c for c in _COLUNAS_OBRIGATORIAS if c not in posicao]
    if faltando:
        raise ListaInvalida(
            "Não achei a(s) coluna(s) obrigatória(s): " + ", ".join(faltando)
            + ". A primeira linha da planilha precisa ser o cabeçalho, com pelo menos "
              "\"cnpj\" e \"razao_social\"."
        )

    def pegar(linha, campo):
        i = posicao.get(campo)
        return linha[i] if i is not None and i < len(linha) else None

    por_cnpj: dict[str, LinhaLista] = {}
    invalidas = repetidas = 0
    for linha in tabela[1:]:
        if not any(c not in (None, "") for c in linha):
            continue  # linha totalmente vazia
        cnpj = _cnpj(pegar(linha, "cnpj"))
        nome = _texto(pegar(linha, "razao_social"))
        if not cnpj or not nome:
            invalidas += 1
            continue
        if cnpj in por_cnpj:
            repetidas += 1  # vale a última ocorrência

        tipo_log, log = _texto(pegar(linha, "tipo_logradouro")), _texto(pegar(linha, "logradouro"))
        ddd, fone = _texto(pegar(linha, "ddd")), _texto(pegar(linha, "telefone"))
        # se o telefone já vem completo (com DDD), não repete o DDD na frente
        telefone = fone if (not ddd or len(re.sub(r"\D", "", fone)) >= 10) else ddd + fone
        por_cnpj[cnpj] = LinhaLista(
            cnpj=cnpj,
            razao_social=nome[:200],
            capital_social=_numero(pegar(linha, "capital_social")),
            porte=_texto(pegar(linha, "porte")),
            situacao=_texto(pegar(linha, "situacao")).upper(),
            cnae=_texto(pegar(linha, "cnae")),
            logradouro=(f"{tipo_log} {log}".strip())[:200],
            numero=_texto(pegar(linha, "numero"))[:20],
            bairro=_texto(pegar(linha, "bairro"))[:100],
            cep=re.sub(r"\D", "", _texto(pegar(linha, "cep")))[:8],
            cidade=_texto(pegar(linha, "cidade")).upper()[:100],
            uf=_texto(pegar(linha, "uf")).upper()[:2],
            telefone=telefone[:30],
            email=_texto(pegar(linha, "email")).lower()[:150],
        )

    if not por_cnpj:
        raise ListaInvalida("Não encontrei nenhuma empresa com CNPJ válido no arquivo.")
    return ListaLida(linhas=list(por_cnpj.values()), invalidas=invalidas, repetidas=repetidas)
